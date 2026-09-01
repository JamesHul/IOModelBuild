"""
Loader for the industry split-out tables.

    python scripts/load_splitout.py

Reads "Piggy IO tables and multipliers_V2 industry splitout tables.xlsx" and
writes build/splitout.pkl. This is the ONLY file that knows that workbook's
layout - the same rule load_sources.py follows for the main supplied file.

What the split-out tables are
-----------------------------
One sheet per (region, model, measure), e.g. 'NSW open va@mp FY23'. Each sheet
is the multiplier for that measure decomposed by the industry that RECEIVES the
impact, rather than collapsed to a single number:

    columns C..DV   FOR USE       - the industry the dollar is spent in (j)
    rows    13..136 FROM INDUSTRY - the industry the impact lands on   (i)

    cell(i, j) = impact on industry i per $1m of final demand in industry j

so a column sums, down the 124 rows, to the multiplier the main model already
uses. Verified against 'Piggy IO tables and multipliers_V2.xlsx': the column sum
equals that region's Simple multiplier for the same measure, 114/114 codes in
all nine regions, to the last bit. check_industry_split.py re-runs that gate.

'open' means the Simple multiplier - initial plus production-induced. It carries
no consumption-induced (induced) effect; the closed sheets do, and are not
loaded yet.

The grid, established by inspection and identical on all 71 sheets
-------------------------------------------------------------------
    r9              blank (single spaces in col C+)
    r10   col B 'FOR USE'         col C+ column IOIG codes
    r11   col B region name       col C+ column industry names
    r12   col A 'FROM INDUSTRY'
    r13-136   114 real codes then 10 'Dummy' rows 9901-9910
    r137  P1 Compensation of employees   - label only, no data
    r139  Sum                            - equals the column sum of r13:r136
    r141-147  Initial effect, First round effect, Simple multiplier,
              Industrial support effect, Production-induced effect,
              Type 1A multiplier, Type 1B multiplier

    col A  IOIG code (TEXT for real codes, NUMBER for the dummies)
    col B  industry name
    col C..DV  the 124-column matrix
    col DW     Final Consumption Expenditure - Households (closed model)
    col DY..EM initial-effect block by measure

Rows 1-8 and 148-274 are empty on every sheet, and are not carried.

Rule 1 of CLAUDE.md applies: the block goes into build/splitout.pkl exactly as
it appears, every row and every column, nothing trimmed or re-spined. The
workbook's MAP layer resolves position with MATCH.
"""
import pickle
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SUP_DIR = ROOT / 'data' / 'supplied'
OUT = ROOT / 'build' / 'splitout.pkl'

SPLIT_FILE = 'Piggy IO tables and multipliers_V2 industry splitout tables.xlsx'

# Model region code -> sheet-name prefix in the split-out file. The prefix is
# 'Qld' where the model says 'QLD'; matching them by string would silently drop
# Queensland, which is exactly the kind of missing region CLAUDE.md rule 5 is
# about.
SHEET_REGION = {'Aus': 'Aus', 'NSW': 'NSW', 'Vic': 'Vic', 'QLD': 'Qld', 'SA': 'SA',
                'WA': 'WA', 'Tas': 'Tas', 'NT': 'NT', 'ACT': 'ACT'}
REGIONS = list(SHEET_REGION)

# (key, sheet suffix, label, unit, number format, matching block in the main
#  supplied multiplier summary). Only the two the model reports today are
# loaded; 'open output', 'open income' and the four 'closed' families use the
# same grid and can be added here without touching anything else.
#
# NOTE ON THE VALUE-ADDED DEFINITION. The split-out file supplies va@mp only -
# value added at MARKET prices, which includes taxes on products (P3). The main
# model's headline is 'Value added multipliers', value added at BASIC prices.
# The two differ; see CLAUDE.md 'GVA has three definitions'. The workbook ties
# to the market-prices block, because that is the only one split by industry,
# and QA_Checks shows the gap against basic prices rather than hiding it.
MEASURES = [
    ('VAmp', 'open va@mp', 'Value added at market prices, $m', '$m', '#,##0.0',
     'Value added at market prices multipliers'),
    ('Employed', 'open employed', 'Employment, FTE', 'FTE', '#,##0.0',
     'Employed multipliers'),
]

FIRST_ROW, LAST_ROW = 9, 147            # the whole non-empty block
LAST_COL = 143                          # through col EM
CODE_ROW, NAME_ROW = 10, 11             # column headers
FIRST_DATA_ROW, LAST_DATA_ROW = 13, 136  # 114 real codes + 10 dummies
FIRST_MAT_COL, LAST_MAT_COL = 3, 126     # col C..DV, same 124 in column order
SUM_ROW = 139
EFFECT_ROWS = (141, 147)

CODE_RE = re.compile(r'^\d{3,4}$')


def norm_code(v):
    """IOIG code as 4-char text. The dummy codes 9901-9910 are stored as numbers
    on these sheets while the real codes are text, so normalise every read."""
    if v is None:
        return None
    s = str(v).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s.zfill(4) if CODE_RE.match(s) else None


def load_block(ws):
    """One sheet, verbatim, plus the indexes the MAP layer needs to find its way
    around without anyone editing the block."""
    rows = [r for r in ws.iter_rows(min_row=FIRST_ROW, max_row=LAST_ROW,
                                    max_col=LAST_COL, values_only=True)]
    verbatim = [list(r) for r in rows]

    def src(r):                       # source row number -> index into verbatim
        return r - FIRST_ROW

    col_codes, col_names = [], []
    for ci in range(FIRST_MAT_COL, LAST_MAT_COL + 1):
        col_codes.append(norm_code(verbatim[src(CODE_ROW)][ci - 1]))
        col_names.append(verbatim[src(NAME_ROW)][ci - 1])

    row_index, row_meta = {}, []
    for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        code = norm_code(verbatim[src(r)][0])
        label = verbatim[src(r)][1]
        row_meta.append({'src_row': r, 'code': code, 'label': label,
                         'dummy': bool(code and code.startswith('99'))})
        if code:
            row_index.setdefault(code, r)

    effects = {}
    for r in range(EFFECT_ROWS[0], EFFECT_ROWS[1] + 1):
        label = verbatim[src(r)][1]
        if label:
            effects[str(label).strip()] = r

    return {'verbatim': verbatim, 'first_row': FIRST_ROW, 'last_row': LAST_ROW,
            'last_col': LAST_COL, 'col_codes': col_codes, 'col_names': col_names,
            'row_index': row_index, 'row_meta': row_meta, 'effect_rows': effects,
            'sum_row': SUM_ROW, 'region_name': verbatim[src(NAME_ROW)][1]}


def main():
    path = SUP_DIR / SPLIT_FILE
    if not path.exists():
        cand = sorted(SUP_DIR.glob('*industry splitout*.xls*'))
        if not cand:
            raise SystemExit(f'not found: {path}')
        path = cand[0]
    print(f'Split-out tables: {path.name}')
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    blocks, missing = {}, []
    for key, suffix, label, unit, fmt, mult_block in MEASURES:
        for region, prefix in SHEET_REGION.items():
            sheet = f'{prefix} {suffix} FY23'
            if sheet not in wb.sheetnames:
                # CLAUDE.md rule 5: never fill a missing region from another.
                # Leave it out and let the QA gate block the run.
                missing.append(sheet)
                print(f'  ! missing sheet {sheet!r} - {region} {key} left EMPTY')
                continue
            d = load_block(wb[sheet])
            d['source'] = f'{path.name} :: {sheet}'
            d['sheet'] = sheet
            blocks[(key, region)] = d
        got = sum(1 for r in REGIONS if (key, r) in blocks)
        print(f'  {key:9s} {got}/{len(REGIONS)} regions, '
              f'{LAST_ROW - FIRST_ROW + 1} rows x {LAST_COL} cols each')
    wb.close()

    # Region distinctness. Six of eight state blocks in the old master model
    # were byte-identical to NSW and nothing in the workbook revealed it.
    import hashlib
    dupes = []
    for key, *_ in MEASURES:
        seen = {}
        for region in REGIONS:
            d = blocks.get((key, region))
            if not d:
                continue
            mat = [row[FIRST_MAT_COL - 1:LAST_MAT_COL]
                   for row in d['verbatim'][FIRST_DATA_ROW - FIRST_ROW:
                                            LAST_DATA_ROW - FIRST_ROW + 1]]
            h = hashlib.sha256(repr(mat).encode()).hexdigest()
            seen.setdefault(h, []).append(region)
        for h, rs in seen.items():
            if len(rs) > 1:
                dupes.append((key, rs))
                print(f'  ! {key}: identical matrices for {rs}')
    if not dupes:
        print('  all region matrices distinct')

    OUT.parent.mkdir(exist_ok=True)
    with open(OUT, 'wb') as fh:
        pickle.dump({'blocks': blocks, 'measures': MEASURES, 'regions': REGIONS,
                     'sheet_region': SHEET_REGION, 'source': path.name,
                     'first_row': FIRST_ROW, 'last_row': LAST_ROW,
                     'last_col': LAST_COL, 'first_data_row': FIRST_DATA_ROW,
                     'last_data_row': LAST_DATA_ROW,
                     'first_mat_col': FIRST_MAT_COL, 'last_mat_col': LAST_MAT_COL,
                     'code_row': CODE_ROW, 'name_row': NAME_ROW,
                     'sum_row': SUM_ROW, 'missing': missing}, fh)
    print(f'\nwrote {OUT.relative_to(ROOT)}')
    return 1 if missing or dupes else 0


if __name__ == '__main__':
    sys.exit(main())
