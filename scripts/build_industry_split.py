"""
Industry-split add-on model.

    python scripts/build_industry_split.py        -> output/IO_Industry_Split_Model.xlsx

A SEPARATE workbook, used only when the question is "which industries does the
impact land in". It does not replace IO_Impact_Model_MASTER.xlsx and it does not
re-derive anything the master already does.

What it takes
-------------
One input: the master model's CALC_Vector - the direct domestic shock AFTER the
margin split, one row per shock line plus one row per (region, margin type)
against the industry that earns the margin. Copy CALC_Vector A6:M<end> from the
master and paste-special VALUES into IN_DirectVector. Nothing else is pasted.

What it does
------------
  IN_DirectVector          the paste, verbatim
      -> CALC_ShockByIndustry   collapses it to one row per (region, industry)
      -> CALC_ShockWide         the same numbers laid out one row per
                                (region, year), 124 columns in the split-out
                                file's own column order, so the multiply below
                                is a plain row-by-row SUMPRODUCT
      -> CALC_IndustryImpact    shock x the industry split-out matrix
      -> OUT_ByIndustry / OUT_ByDivision / OUT_Matrix

The multiply, precisely. For region r, year y and receiving industry i:

    Total(i)    = SUM over j of  Shock(r, j, y) * M_r[i][j]
    Direct(i)   = Shock(r, i, y) * InitialEffect_r[i]
    Indirect(i) = Total(i) - Direct(i)

M_r is the region's split-out matrix, held verbatim on RAW_Split_*. Its column
j sums, down the 124 rows, to that region's Simple multiplier for industry j -
which is exactly Direct + Indirect in the master model. The initial effect lands
wholly in the industry that was shocked, so subtracting it off the diagonal is
what separates direct from indirect. QA_Reconcile re-derives both totals
straight off the multiplier rows of the same RAW block and reports the
difference, so "does this sum to the multipliers the main model uses" is a
number in the workbook, not a claim in a README.

Measures
--------
Today: 'open va@mp' and 'open employed'. 'open' is the Simple multiplier -
initial plus production-induced, no consumption-induced (induced) effect. The
supplied file also carries open output/income and four closed families on an
identical grid; add them to MEASURES in scripts/load_splitout.py and rebuild.

READ THIS ABOUT VALUE ADDED. The split-out file supplies va@mp - value added at
MARKET prices, which includes taxes on products (P3). The master model's
headline is value added at BASIC prices. They are different definitions and they
do not tie; the gap is about +0.01 per dollar on average nationally, and is
signed both ways by industry. QA_Checks prints it. Do not present the industry
split as a decomposition of the master's headline GVA without saying which
definition it is on.

Environment
-----------
  IOSPLIT_YEARS=10    year columns; must match the master's CALC_Vector width
  IOSPLIT_LINES=700   rows on the IN_DirectVector paste area
  IOSPLIT_EXAMPLE=0   build with an empty paste area instead of the worked example
  IOSPLIT_SUBSET=1    2 regions, 2 years - small enough to recalculate in-process
"""
import os
import pickle
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter as CL

ROOT = Path(__file__).resolve().parent.parent
SPLITPKL = ROOT / 'build' / 'splitout.pkl'
OUTFILE = ROOT / 'output' / 'IO_Industry_Split_Model.xlsx'
SUBSET_OUT = ROOT / 'build' / 'IO_Industry_Split_subset.xlsx'

H1 = Font(bold=True, size=14, color='FFFFFF')
H2 = Font(bold=True, size=11)
H3 = Font(bold=True, size=10)
BLUE = Font(color='0000CC')
GREEN = Font(color='006100')
NOTE = Font(italic=True, size=9, color='666666')
WARN = Font(bold=True, color='9C0006')
TITLEFILL = PatternFill('solid', fgColor='1F3864')
HDRFILL = PatternFill('solid', fgColor='DDEBF7')
GREYFILL = PatternFill('solid', fgColor='F2F2F2')
YELLOW = PatternFill('solid', fgColor='FFF2CC')
MONEY = '#,##0.0'
NUM6 = '0.000000'
PCT = '0.0%'
NUM4 = '0.0000'

NYR = int(os.environ.get('IOSPLIT_YEARS', '10'))
NIN = int(os.environ.get('IOSPLIT_LINES', '700'))
EXAMPLE = os.environ.get('IOSPLIT_EXAMPLE', '1') != '0'
SUBSET = bool(os.environ.get('IOSPLIT_SUBSET'))

# ANZSIC 2006 division for an IOIG(2022) code, keyed on the two-digit
# subdivision prefix the IOIG code carries. All nineteen divisions A-S are
# represented on the 114-code spine - there is no eighteen-division ANZSIC
# grouping, so if a study needs eighteen lines, merge two rows by editing the
# Lists tab and every OUT_ tab follows.
DIVISIONS = [
    ('A', 'Agriculture, Forestry and Fishing', ('01', '02', '03', '04', '05')),
    ('B', 'Mining', ('06', '07', '08', '09', '10')),
    ('C', 'Manufacturing', tuple(f'{n:02d}' for n in range(11, 26))),
    ('D', 'Electricity, Gas, Water and Waste Services', ('26', '27', '28', '29')),
    ('E', 'Construction', ('30', '31', '32')),
    ('F', 'Wholesale Trade', ('33', '34', '35', '36', '37', '38')),
    ('G', 'Retail Trade', ('39', '40', '41', '42', '43')),
    ('H', 'Accommodation and Food Services', ('44', '45')),
    ('I', 'Transport, Postal and Warehousing',
     ('46', '47', '48', '49', '50', '51', '52', '53')),
    ('J', 'Information Media and Telecommunications',
     ('54', '55', '56', '57', '58', '59', '60')),
    ('K', 'Financial and Insurance Services', ('62', '63', '64')),
    ('L', 'Rental, Hiring and Real Estate Services', ('66', '67')),
    ('M', 'Professional, Scientific and Technical Services', ('69', '70')),
    ('N', 'Administrative and Support Services', ('72', '73')),
    ('O', 'Public Administration and Safety', ('75', '76', '77')),
    ('P', 'Education and Training', ('80', '81', '82')),
    ('Q', 'Health Care and Social Assistance', ('84', '85', '86', '87')),
    ('R', 'Arts and Recreation Services', ('89', '90', '91', '92')),
    ('S', 'Other Services', ('94', '95', '96')),
]
DIV_OF_PREFIX = {p: (dc, dn) for dc, dn, ps in DIVISIONS for p in ps}
DUMMY_DIV = ('Z', 'Dummy / unallocated')

# 6700 has no row in the supplied 114-code set; the master already bridges it on
# MAP_ShockKeys, but a hand-built paste might not, so the bridge is repeated
# here and shown on the Lists tab.
BRIDGE = [('6700', '6701',
           'Imputed rent for owner-occupiers. The supplied 114-code spine has no '
           '6700 and treats ownership of dwellings as a single 6701.')]

# A worked example so the chain is alive when the file is opened. Every one of
# these rows is blue - hardcoded input - and IN_DirectVector says so at the top.
EXAMPLE_ROWS = [
    ('line',   'NSW', '3101', 100.0), ('line',   'NSW', '3201',  40.0),
    ('line',   'Vic', '1101',  20.0), ('line',   'QLD', '6901',  15.0),
    ('line',   'Aus', '3101',  50.0), ('margin', 'NSW', '3301',   8.0),
    ('margin', 'NSW', '3901',   5.0), ('margin', 'Vic', '4601',   1.2),
]


def band(ws, row, text, width):
    ws.cell(row=row, column=1, value=text).font = H1
    for c in range(1, width + 1):
        ws.cell(row=row, column=c).fill = TITLEFILL


def note(ws, row, text, col=1):
    ws.cell(row=row, column=col, value=text).font = NOTE


def hdr(ws, row, labels, col0=1):
    for i, h in enumerate(labels, col0):
        c = ws.cell(row=row, column=i, value=h)
        c.font = H3
        c.fill = HDRFILL
        c.alignment = Alignment(wrap_text=True, vertical='bottom')


def whole_row(rng, r, ncols):
    """A whole row of `rng` as a RANGE, not an array.

    INDEX(rng, r, 0) is the array form and Excel understands it, but it is an
    array where a range is wanted and the `formulas` engine evaluates it as the
    first cell alone - which would have made the core multiply silently wrong
    and unverifiable. INDEX:INDEX returns a real range reference and behaves the
    same in Excel and in the engine, so the recalculation actually checks the
    arithmetic. Neither INDEX form is a banned function.
    """
    # Parenthesised: ':' binds tighter than '*' in Excel so it is already
    # correct bare, but CLAUDE.md's precedence rule says do not rely on that.
    return f'(INDEX({rng},{r},1):INDEX({rng},{r},{ncols}))'


def division_of(code):
    if code is None:
        return ('', '')
    if code.startswith('99'):
        return DUMMY_DIV
    # Sentinel deliberately a word: '?' is a single-character wildcard in a
    # COUNTIF criterion, so a gate looking for it would match every division.
    return DIV_OF_PREFIX.get(code[:2], ('UNMAPPED', 'UNMAPPED - add the prefix to DIVISIONS'))


def main():
    if not SPLITPKL.exists():
        raise SystemExit('build/splitout.pkl missing - run scripts/load_splitout.py first')
    src = pickle.load(open(SPLITPKL, 'rb'))
    blocks, regions = src['blocks'], list(src['regions'])
    measures = [m for m in src['measures']]
    global NIN
    if SUBSET:
        # Structurally identical, just small enough for the `formulas` engine to
        # load and calculate in one sitting. Never a delivery build.
        regions = regions[:2]
        NIN = 40
    nyr = 2 if SUBSET else NYR
    nreg = len(regions)

    # The 124-entry column order, taken from the split-out file itself. Rows and
    # columns share it, which is what lets a coefficient row and a shock row be
    # multiplied position by position.
    ref = blocks[(measures[0][0], regions[0])]
    order = [(m['code'], m['label']) for m in ref['row_meta']]
    NC = len(order)
    for key, _s, _l, _u, _f, _b in measures:
        for rg in regions:
            b = blocks.get((key, rg))
            if b is None:
                raise SystemExit(f'{key}/{rg} missing from splitout.pkl - '
                                 'never fill a region from another; fix the source')
            if [m['code'] for m in b['row_meta']] != [c for c, _ in order]:
                raise SystemExit(f'{key}/{rg} row order differs from {regions[0]}')
            if b['col_codes'] != [c for c, _ in order]:
                raise SystemExit(f'{key}/{rg} column order differs from its own row order')
    region_name = {rg: str(blocks[(measures[0][0], rg)]['region_name']) for rg in regions}

    FR, LR, LC = src['first_row'], src['last_row'], src['last_col']
    NSRC = LR - FR + 1                       # 139 source rows per region block
    NRAWCOL = 6 + LC                         # 6 key columns then the source block
    MATC0, MATC1 = 6 + src['first_mat_col'], 6 + src['last_mat_col']   # I .. EB
    CODEC, NAMEC = 7, 8                      # source col A and B, once shifted

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ============================================================ README
    ws = wb.create_sheet('README')
    band(ws, 1, 'Industry split of the IO impact - add-on to IO_Impact_Model_MASTER', 8)
    lines = [
        ('What this is',
         'A separate workbook that answers "which industries does the impact land in". '
         'It takes the master model\'s post-margin shock table as an input and multiplies '
         'it through the supplied industry split-out matrices. It does not re-derive '
         'anything the master already does, and it is not a replacement for it.'),
        ('The one input',
         'IN_DirectVector. Open IO_Impact_Model_MASTER.xlsx, let it recalculate, then copy '
         f'CALC_Vector A6:{CL(3 + NYR)}<last row> and paste-special VALUES into '
         'IN_DirectVector A6. That block is the direct domestic shock AFTER the margin '
         'split: one row per shock line, plus one row per region and margin type against '
         'the industry that earns the margin.'),
        ('Measures today',
         'Open value added at market prices, and open employment (FTE). "Open" is the '
         'Simple multiplier - initial plus production-induced. No induced effect. The '
         'supplied file carries open output/income and four closed families on the same '
         'grid; add them to MEASURES in scripts/load_splitout.py and rebuild.'),
        ('Value added definition - READ THIS',
         'The split-out file supplies value added at MARKET prices only. The master '
         'model\'s headline is value added at BASIC prices. These are different '
         'definitions and they will not tie. QA_Checks shows the gap. Label any output '
         'from this workbook "value added at market prices".'),
        ('Does it tie to the main model',
         'Yes, by construction and it is checked. Each column of a split-out matrix sums '
         'to that region\'s Simple multiplier for the same measure - verified against '
         'Piggy IO tables and multipliers_V2.xlsx, 114/114 codes in all nine regions, to '
         'the last bit. QA_Reconcile re-derives the total and the direct effect straight '
         'off the multiplier rows of the same RAW block and reports the difference.'),
        ('The chain',
         'IN_DirectVector -> CALC_ShockByIndustry -> CALC_ShockWide -> CALC_IndustryImpact '
         '-> OUT_ByIndustry / OUT_ByDivision / OUT_Matrix. RAW_Split_* hold the supplied '
         'matrices verbatim; MAP_SplitIndex resolves every position with MATCH.'),
        ('Summary industries',
         f'ANZSIC 2006 divisions A-S. All nineteen are represented on the 114-code spine, '
         'so the summary is nineteen lines, not eighteen. The map is a visible table on '
         'Lists - merge two rows there if a study needs a different grouping and every '
         'OUT_ tab follows.'),
        ('Colour key',
         'Blue text = hardcoded input. Black = formula. Green = cross-sheet link. '
         'Yellow fill = the analyst completes it.'),
        ('Rebuild',
         'python scripts/load_splitout.py && python scripts/build_industry_split.py. '
         'Never hand-edit this workbook.'),
    ]
    for i, (a, b) in enumerate(lines):
        r = 3 + i * 2
        ws.cell(row=r, column=1, value=a).font = H2
        c = ws.cell(row=r + 1, column=1, value=b)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if 'READ THIS' in a:
            c.font = WARN
    ws.column_dimensions['A'].width = 150
    for i in range(len(lines)):
        ws.row_dimensions[4 + i * 2].height = 46

    # ============================================================ Settings
    ws = wb.create_sheet('Settings')
    band(ws, 1, 'Settings', 6)
    rows = [
        ('Built', str(date.today())),
        ('Split-out source', src['source']),
        ('Vintage of the split-out tables', 'FY23 (2022-23), same as the supplied '
                                            'multiplier set'),
        ('Year columns', nyr),
        ('Paste rows on IN_DirectVector', NIN),
        ('Regions', ', '.join(regions)),
        ('Measures', ', '.join(m[2] for m in measures)),
        ('Spine', f'{sum(1 for c, _ in order if not c.startswith("99"))} real IOIG codes '
                  f'+ {sum(1 for c, _ in order if c.startswith("99"))} dummy rows '
                  f'= {NC} columns, in the split-out file\'s own order'),
        ('Build', 'SUBSET - test harness, not a delivery build' if SUBSET else 'FULL'),
    ]
    hdr(ws, 4, ['Setting', 'Value'])
    for i, (a, b) in enumerate(rows):
        ws.cell(row=5 + i, column=1, value=a).font = H3
        ws.cell(row=5 + i, column=2, value=b).font = BLUE
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 90

    # ============================================================ Lists
    ws = wb.create_sheet('Lists')
    band(ws, 1, 'Lists - regions, the 6700 bridge, ANZSIC divisions, measures', 8)
    hdr(ws, 5, ['#', 'Region (model code)', 'Region name', 'Sheet prefix in the '
                                                           'split-out file'])
    LR0 = 6
    for i, rg in enumerate(regions):
        ws.cell(row=LR0 + i, column=1, value=i + 1)
        ws.cell(row=LR0 + i, column=2, value=rg).font = BLUE
        ws.cell(row=LR0 + i, column=3, value=region_name[rg]).font = BLUE
        ws.cell(row=LR0 + i, column=4,
                value=src['sheet_region'][rg]).font = BLUE
    LR1 = LR0 + nreg - 1

    BR_T = LR1 + 2
    ws.cell(row=BR_T, column=1, value='Code bridge - model spine to the supplied '
                                      '114-code set').font = H2
    hdr(ws, BR_T + 1, ['#', 'Model code', 'Code used here', 'Why'])
    BR0 = BR_T + 2
    for i, (a, b, why) in enumerate(BRIDGE):
        ws.cell(row=BR0 + i, column=1, value=i + 1)
        c = ws.cell(row=BR0 + i, column=2, value=a)
        c.number_format = '@'
        c.font = BLUE
        c = ws.cell(row=BR0 + i, column=3, value=b)
        c.number_format = '@'
        c.font = BLUE
        ws.cell(row=BR0 + i, column=4, value=why).font = NOTE
    BR1 = BR0 + len(BRIDGE) - 1

    DV_T = BR1 + 2
    ws.cell(row=DV_T, column=1, value='Summary industries - ANZSIC 2006 divisions').font = H2
    note(ws, DV_T + 1, 'All nineteen divisions are represented on the 114-code spine. To '
                       'report a different grouping, edit the Division name here; every '
                       'OUT_ tab reads it and follows.')
    hdr(ws, DV_T + 2, ['#', 'Division', 'Division name', 'IOIG codes on the spine'])
    DV0 = DV_T + 3
    div_list = [(dc, dn) for dc, dn, _ in DIVISIONS] + [DUMMY_DIV]
    for i, (dc, dn) in enumerate(div_list):
        ws.cell(row=DV0 + i, column=1, value=i + 1)
        ws.cell(row=DV0 + i, column=2, value=dc).font = BLUE
        ws.cell(row=DV0 + i, column=3, value=dn).font = BLUE
        ws.cell(row=DV0 + i, column=4,
                value=sum(1 for c, _ in order if division_of(c)[0] == dc))
    DV1 = DV0 + len(div_list) - 1

    MS_T = DV1 + 2
    ws.cell(row=MS_T, column=1, value='Measures built into this workbook').font = H2
    hdr(ws, MS_T + 1, ['#', 'Key', 'Label', 'Unit', 'Split-out sheets',
                       'Matching block in the supplied multiplier summary'])
    MS0 = MS_T + 2
    for i, (key, suffix, label, unit, fmt, mblock) in enumerate(measures):
        ws.cell(row=MS0 + i, column=1, value=i + 1)
        ws.cell(row=MS0 + i, column=2, value=key).font = BLUE
        ws.cell(row=MS0 + i, column=3, value=label).font = BLUE
        ws.cell(row=MS0 + i, column=4, value=unit).font = BLUE
        ws.cell(row=MS0 + i, column=5, value=f'<region> {suffix} FY23').font = BLUE
        ws.cell(row=MS0 + i, column=6, value=mblock).font = BLUE
    for col, w in zip('ABCDEF', [5, 22, 46, 8, 26, 46]):
        ws.column_dimensions[col].width = w

    # ============================================================ Lists_Spine
    ws = wb.create_sheet('Lists_Spine')
    band(ws, 1, 'Lists_Spine - the 124-entry order the split-out tables use for both '
                'their rows and their columns', 6)
    note(ws, 2, 'Position is what makes the multiply work: coefficient column j and shock '
                'column j are the same j. Taken from the split-out file, never retyped.')
    hdr(ws, 5, ['#', 'IOIG code', 'Industry name', 'Division', 'Division name', 'Real code?'])
    SP0 = 6
    for i, (code, label) in enumerate(order):
        r = SP0 + i
        ws.cell(row=r, column=1, value=i + 1)
        c = ws.cell(row=r, column=2, value=code)
        c.number_format = '@'
        c.font = BLUE
        ws.cell(row=r, column=3, value=label).font = BLUE
        dc, dn = division_of(code)
        ws.cell(row=r, column=4, value=dc).font = BLUE
        ws.cell(row=r, column=5, value=dn).font = BLUE
        ws.cell(row=r, column=6, value='' if code.startswith('99') else 'Y')
    SP1 = SP0 + NC - 1
    ws.column_dimensions['B'].width = 11
    ws.column_dimensions['C'].width = 62
    ws.column_dimensions['E'].width = 44
    ws.freeze_panes = 'A6'

    # ============================================================ RAW_Split_*
    raw_tab, raw_last = {}, {}
    for key, suffix, label, unit, fmt, mblock in measures:
        tab = f'RAW_Split_{key}'
        raw_tab[key] = tab
        ws = wb.create_sheet(tab)
        band(ws, 1, f'{tab} - the supplied "{suffix}" split-out matrices, VERBATIM, '
                    'stacked with a jurisdiction column', 12)
        note(ws, 2, f'Source: {src["source"]}, sheets "<region> {suffix} FY23". Source rows '
                    f'{FR}-{LR} and columns A-{CL(LC)} exactly as supplied - every row '
                    'including the dummies, the Sum row and the seven effect rows, every '
                    'column including the closed-model and initial-effect blocks. Nothing '
                    'trimmed, nothing re-spined.')
        note(ws, 3, 'Columns A-F are the stack key, added here and nowhere in the source: '
                    'they are what lets nine sheets live on one tab. Column G onwards is '
                    f'source column A onwards; the {NC}-column matrix is '
                    f'{CL(MATC0)}:{CL(MATC1)}.')
        hdr(ws, 5, ['Key (Region|Code or Region|Label)', 'Region', 'Code', 'Row type',
                    'Source sheet', 'Src row']
            + [f'src {CL(c)}' for c in range(1, LC + 1)])
        R0 = 6
        r = R0
        for rg in regions:
            b = blocks[(key, rg)]
            meta_by_srcrow = {m['src_row']: m for m in b['row_meta']}
            for si, row in enumerate(b['verbatim']):
                srow = FR + si
                m = meta_by_srcrow.get(srow)
                if m is not None:
                    code = m['code']
                    rtype = 'Dummy' if m['dummy'] else 'Spine'
                    lab = code
                elif srow in (src['code_row'], src['name_row']):
                    code, rtype = None, 'Header'
                    lab = 'Codes' if srow == src['code_row'] else 'Names'
                elif srow == src['sum_row']:
                    code, rtype, lab = None, 'Sum', 'Sum'
                else:
                    lbl = row[1]
                    lbl = str(lbl).strip() if lbl is not None else ''
                    code = None
                    if lbl in b['effect_rows']:
                        rtype, lab = 'Effect', lbl
                    elif lbl:
                        rtype, lab = 'Other', lbl
                    else:
                        rtype, lab = 'Blank', ''
                c = ws.cell(row=r, column=1, value=f'{rg}|{lab}' if lab else '')
                c.number_format = '@'
                ws.cell(row=r, column=2, value=rg)
                c = ws.cell(row=r, column=3, value=code)
                c.number_format = '@'
                ws.cell(row=r, column=4, value=rtype)
                ws.cell(row=r, column=5, value=b['sheet'])
                ws.cell(row=r, column=6, value=srow)
                for ci, v in enumerate(row, 7):
                    if v is None:
                        continue
                    cc = ws.cell(row=r, column=ci, value=v)
                    if ci in (CODEC,) and isinstance(v, str):
                        cc.number_format = '@'
                r += 1
        R1 = r - 1
        raw_last[key] = R1
        ws.freeze_panes = 'G6'
        for col, w in zip(['A', 'B', 'C', 'D', 'E', 'F'], [22, 8, 8, 10, 26, 8]):
            ws.column_dimensions[col].width = w

    RK = {k: f"{raw_tab[k]}!$A${6}:$A${raw_last[k]}" for k in raw_tab}
    RC = {k: f"{raw_tab[k]}!$C${6}:$C${raw_last[k]}" for k in raw_tab}
    MAT = {k: f"{raw_tab[k]}!${CL(MATC0)}${6}:${CL(MATC1)}${raw_last[k]}" for k in raw_tab}

    # ============================================================ IN_DirectVector
    ws = wb.create_sheet('IN_DirectVector')
    band(ws, 1, 'IN_DirectVector - PASTE THE MASTER MODEL\'S CALC_Vector HERE', 8 + NYR)
    ws.cell(row=2, column=1, value=(
        f'Open IO_Impact_Model_MASTER.xlsx, let it recalculate, copy CALC_Vector '
        f'A6:{CL(3 + nyr)}<last row> and paste-special VALUES into A6 below. That block is '
        'the direct domestic shock after the margin split - shock lines plus the '
        'reallocated margins, each against the industry that earns it.')).font = WARN
    note(ws, 3, 'Columns A-' + CL(3 + nyr) + ' are the paste. Columns ' + CL(4 + nyr) +
                ' onward are formulas that check it and build the lookup key; do not '
                'paste over them.')
    IVK = 4 + nyr           # Use code
    IVR = IVK + 1           # Region OK
    IVC = IVK + 2           # Code OK
    IVKEY = IVK + 3         # Key
    IVT = IVK + 4           # Line total
    hdr(ws, 5, ['Source', 'Region', 'Code'] + [f'Year {y + 1}' for y in range(nyr)]
        + ['Use code (6700 bridged)', 'Region check', 'Code check',
           'Key (Region|Use code)', 'Line total'])
    IV0 = 6
    IV1 = IV0 + NIN - 1
    for i in range(NIN):
        r = IV0 + i
        for c in range(1, IVK):
            cell = ws.cell(row=r, column=c)
            cell.fill = YELLOW
            cell.font = BLUE
            if c == 3:
                cell.number_format = '@'
            elif c >= 4:
                cell.number_format = MONEY
        c = ws.cell(row=r, column=IVK, value=(
            f'=IF($C{r}="","",IFERROR(INDEX(Lists!$C${BR0}:$C${BR1},'
            f'MATCH($C{r},Lists!$B${BR0}:$B${BR1},0)),$C{r}))'))
        c.number_format = '@'
        ws.cell(row=r, column=IVR, value=(
            f'=IF($B{r}="","",IF(ISNUMBER(MATCH($B{r},Lists!$B${LR0}:$B${LR1},0)),'
            f'"OK","UNKNOWN REGION"))'))
        ws.cell(row=r, column=IVC, value=(
            f'=IF(${CL(IVK)}{r}="","",IF(ISNUMBER(MATCH(${CL(IVK)}{r},'
            f'Lists_Spine!$B${SP0}:$B${SP1},0)),"OK","NOT ON THE 114-CODE SPINE"))'))
        c = ws.cell(row=r, column=IVKEY, value=(
            f'=IF(OR($B{r}="",${CL(IVK)}{r}=""),"",$B{r}&"|"&${CL(IVK)}{r})'))
        c.number_format = '@'
        # 0, not "", on an unused row: in Excel ""<>0 is TRUE, so returning text
        # here would make every empty row count as a row with spend.
        ws.cell(row=r, column=IVT, value=(
            f'=IF($B{r}="",0,SUM($D{r}:${CL(3 + nyr)}{r}))')).number_format = MONEY
    if EXAMPLE:
        ws.cell(row=4, column=1, value=(
            'EXAMPLE DATA IS LOADED BELOW so the chain is alive on open. It is NOT a '
            'result. Delete A6:' + CL(3 + nyr) + str(IV0 + len(EXAMPLE_ROWS) - 1) +
            ' and paste your own CALC_Vector block.')).font = WARN
        for i, (srcname, rg, code, amt) in enumerate(EXAMPLE_ROWS):
            if rg not in regions:
                continue
            r = IV0 + i
            ws.cell(row=r, column=1, value=srcname).font = BLUE
            ws.cell(row=r, column=2, value=rg).font = BLUE
            c = ws.cell(row=r, column=3, value=code)
            c.number_format = '@'
            c.font = BLUE
            for y in range(nyr):
                v = amt if y == 0 else round(amt * (0.6 ** y), 4)
                cc = ws.cell(row=r, column=4 + y, value=v)
                cc.number_format = MONEY
                cc.font = BLUE
    ws.freeze_panes = 'D6'
    for col, w in zip(['A', 'B', 'C'], [10, 9, 10]):
        ws.column_dimensions[col].width = w

    IVKEYR = f'IN_DirectVector!${CL(IVKEY)}${IV0}:${CL(IVKEY)}${IV1}'
    IVSRCR = f'IN_DirectVector!$A${IV0}:$A${IV1}'

    # ============================================================ MAP_SplitIndex
    ws = wb.create_sheet('MAP_SplitIndex')
    band(ws, 1, 'MAP_SplitIndex - resolve every position in the RAW blocks once, with '
                'MATCH on region and code', 11)
    note(ws, 2, 'One MATCH per measure, region and industry, plus one for each region\'s '
                'Initial effect and Simple multiplier rows. Everything downstream indexes '
                'off these row numbers instead of repeating the lookup, and the alignment '
                'column proves the row it landed on really carries that code.')
    hdr(ws, 5, ['Measure', 'Region', 'Position', 'IOIG code', 'Industry', 'Division',
                'Key', 'RAW row (matrix)', 'RAW row (Initial effect)',
                'RAW row (Simple multiplier)', 'Alignment'])
    MI0 = 6
    r = MI0
    mi_row = {}
    for key, *_ in measures:
        for rg in regions:
            for pos, (code, lbl) in enumerate(order, 1):
                ws.cell(row=r, column=1, value=key)
                ws.cell(row=r, column=2, value=rg)
                ws.cell(row=r, column=3, value=pos)
                c = ws.cell(row=r, column=4, value=f'=Lists_Spine!$B${SP0 + pos - 1}')
                c.number_format = '@'
                c.font = GREEN
                ws.cell(row=r, column=5,
                        value=f'=Lists_Spine!$C${SP0 + pos - 1}').font = GREEN
                ws.cell(row=r, column=6,
                        value=f'=Lists_Spine!$D${SP0 + pos - 1}').font = GREEN
                c = ws.cell(row=r, column=7, value=f'=$B{r}&"|"&$D{r}')
                c.number_format = '@'
                ws.cell(row=r, column=8,
                        value=f'=IFERROR(MATCH($G{r},{RK[key]},0),0)')
                ws.cell(row=r, column=9, value=(
                    f'=IFERROR(MATCH($B{r}&"|Initial effect",{RK[key]},0),0)'))
                ws.cell(row=r, column=10, value=(
                    f'=IFERROR(MATCH($B{r}&"|Simple multiplier",{RK[key]},0),0)'))
                ws.cell(row=r, column=11, value=(
                    f'=IF($H{r}=0,"NO ROW",IF(EXACT(INDEX({RC[key]},$H{r}),$D{r}),'
                    f'"OK","MISALIGNED"))'))
                mi_row[(key, rg, pos)] = r
                r += 1
    MI1 = r - 1
    ws.freeze_panes = 'D6'
    for col, w in zip(['A', 'B', 'C', 'D', 'E', 'F'], [11, 8, 9, 10, 50, 9]):
        ws.column_dimensions[col].width = w

    # ============================================================ CALC_ShockByIndustry
    ws = wb.create_sheet('CALC_ShockByIndustry')
    band(ws, 1, 'CALC_ShockByIndustry - the shock, post margin allocation, as one row per '
                'region and industry', 10 + nyr)
    note(ws, 2, 'This is the paste collapsed onto the spine: every shock line and every '
                'reallocated margin dollar summed onto the industry that receives it. '
                'SUMIFS on the Region|Code key, so a code stored as text stays text.')
    hdr(ws, 5, ['Region', 'IOIG code', 'Industry', 'Division', 'Key']
        + [f'Year {y + 1}' for y in range(nyr)]
        + ['Total, all years', 'Year 1 from shock lines', 'Year 1 from margins',
           'Margin share, year 1'])
    CS0 = 6
    CSY = 6                                 # first year column
    r = CS0
    for rg in regions:
        for pos, (code, lbl) in enumerate(order, 1):
            ws.cell(row=r, column=1, value=rg)
            c = ws.cell(row=r, column=2, value=f'=Lists_Spine!$B${SP0 + pos - 1}')
            c.number_format = '@'
            c.font = GREEN
            ws.cell(row=r, column=3,
                    value=f'=Lists_Spine!$C${SP0 + pos - 1}').font = GREEN
            ws.cell(row=r, column=4,
                    value=f'=Lists_Spine!$D${SP0 + pos - 1}').font = GREEN
            c = ws.cell(row=r, column=5, value=f'=$A{r}&"|"&$B{r}')
            c.number_format = '@'
            for y in range(nyr):
                ws.cell(row=r, column=CSY + y, value=(
                    f'=SUMIFS(IN_DirectVector!${CL(4 + y)}${IV0}:${CL(4 + y)}${IV1},'
                    f'{IVKEYR},$E{r})')).number_format = MONEY
            ws.cell(row=r, column=CSY + nyr, value=(
                f'=SUM(${CL(CSY)}{r}:${CL(CSY + nyr - 1)}{r})')).number_format = MONEY
            ws.cell(row=r, column=CSY + nyr + 1, value=(
                f'=SUMIFS(IN_DirectVector!$D${IV0}:$D${IV1},{IVKEYR},$E{r},'
                f'{IVSRCR},"line")')).number_format = MONEY
            ws.cell(row=r, column=CSY + nyr + 2, value=(
                f'=SUMIFS(IN_DirectVector!$D${IV0}:$D${IV1},{IVKEYR},$E{r},'
                f'{IVSRCR},"margin")')).number_format = MONEY
            ws.cell(row=r, column=CSY + nyr + 3, value=(
                f'=IFERROR(${CL(CSY + nyr + 2)}{r}/'
                f'(${CL(CSY + nyr + 1)}{r}+${CL(CSY + nyr + 2)}{r}),"")')
            ).number_format = PCT
            r += 1
    CS1 = r - 1
    CSTOT = CSY + nyr
    ws.freeze_panes = 'F6'
    # Deliberately no autofilter. CALC_ShockWide links to these rows by position,
    # so sorting this tab would silently re-point every link. Filter and sort on
    # OUT_ByIndustry instead, which links by row and moves as a unit.
    ws.cell(row=4, column=1, value='DO NOT SORT OR INSERT ROWS - CALC_ShockWide '
                                   'links to these rows by position.').font = WARN
    for col, w in zip(['A', 'B', 'C', 'D', 'E'], [8, 10, 52, 9, 14]):
        ws.column_dimensions[col].width = w
    CSREG = f'CALC_ShockByIndustry!$A${CS0}:$A${CS1}'

    # ============================================================ CALC_ShockWide
    ws = wb.create_sheet('CALC_ShockWide')
    band(ws, 1, 'CALC_ShockWide - the same shock, one row per region and year, laid out in '
                'the split-out column order', 6 + NC)
    note(ws, 2, 'Only a layout change, and only so the multiply on CALC_IndustryImpact is a '
                'plain SUMPRODUCT of two horizontal ranges - a coefficient row against a '
                'shock row. Each cell is a direct link to CALC_ShockByIndustry.')
    note(ws, 3, 'Row 6 proves the link landed on the right industry: it compares the '
                'header code against the code on the row actually linked. Column '
                f'{CL(4 + NC + 1)} proves the row totals against the long table.')
    CWD = 4                                 # first data column
    hdr(ws, 5, ['Region', 'Year', 'Key'])
    for pos, (code, lbl) in enumerate(order, 1):
        c = ws.cell(row=5, column=CWD + pos - 1, value=code)
        c.number_format = '@'
        c.font = H3
        c.fill = HDRFILL
    ws.cell(row=5, column=CWD + NC, value='Row sum').font = H3
    ws.cell(row=5, column=CWD + NC).fill = HDRFILL
    ws.cell(row=5, column=CWD + NC + 1, value='Check vs long table').font = H3
    ws.cell(row=5, column=CWD + NC + 1).fill = HDRFILL
    ws.cell(row=6, column=3, value='Alignment').font = H3
    for pos in range(1, NC + 1):
        ws.cell(row=6, column=CWD + pos - 1, value=(
            f'=IF(EXACT({CL(CWD + pos - 1)}$5,CALC_ShockByIndustry!$B${CS0 + pos - 1}),'
            f'"","X")'))
    ws.cell(row=6, column=CWD + NC + 1, value=(
        f'=IF(COUNTIF(${CL(CWD)}$6:${CL(CWD + NC - 1)}$6,"X")=0,"ALIGNED",'
        f'"COLUMN ORDER BROKEN")')).font = H3
    CW0 = 7
    r = CW0
    for ri, rg in enumerate(regions):
        for y in range(nyr):
            ws.cell(row=r, column=1, value=rg)
            ws.cell(row=r, column=2, value=y + 1)
            c = ws.cell(row=r, column=3, value=f'=$A{r}&"|"&$B{r}')
            c.number_format = '@'
            for pos in range(1, NC + 1):
                ws.cell(row=r, column=CWD + pos - 1, value=(
                    f'=CALC_ShockByIndustry!${CL(CSY + y)}${CS0 + ri * NC + pos - 1}')
                ).number_format = MONEY
            ws.cell(row=r, column=CWD + NC, value=(
                f'=SUM(${CL(CWD)}{r}:${CL(CWD + NC - 1)}{r})')).number_format = MONEY
            # This row's year is fixed by the loop, so the column is a literal
            # rather than an INDEX with a computed column number - simpler to
            # read and it keeps this a plain SUMIFS.
            ws.cell(row=r, column=CWD + NC + 1, value=(
                f'=ROUND(${CL(CWD + NC)}{r}-SUMIFS('
                f'CALC_ShockByIndustry!${CL(CSY + y)}${CS0}:${CL(CSY + y)}${CS1},'
                f'{CSREG},$A{r}),6)')).number_format = NUM6
            r += 1
    CW1 = r - 1
    ws.freeze_panes = 'D7'
    WIDE = f'CALC_ShockWide!${CL(CWD)}${CW0}:${CL(CWD + NC - 1)}${CW1}'
    WIDEKEY = f'CALC_ShockWide!$C${CW0}:$C${CW1}'

    # ============================================================ CALC_IndustryImpact
    ws = wb.create_sheet('CALC_IndustryImpact')
    band(ws, 1, 'CALC_IndustryImpact - shock x split-out matrix, by receiving industry', 14)
    note(ws, 2, 'Total(i) = SUM over j of Shock(j) x M[i][j], the coefficient row for '
                'industry i against the shock row for that region and year. Direct(i) is '
                'the shock into industry i times its own initial effect - the initial '
                'effect lands wholly in the industry that was shocked - and Indirect is '
                'the remainder.')
    cols = ['Measure', 'Region', 'Position', 'IOIG code', 'Industry', 'Division',
            'RAW row', 'RAW row (Initial)']
    for y in range(nyr):
        cols += [f'Yr{y + 1} Direct', f'Yr{y + 1} Indirect', f'Yr{y + 1} Total']
    cols += ['Direct, all years', 'Indirect, all years', 'TOTAL, all years']
    hdr(ws, 5, cols)
    CIY = 9                                  # first year triple
    CIDT = CIY + nyr * 3
    CI0 = 6
    r = CI0
    ci_row = {}
    for key, suffix, label, unit, fmt, mblock in measures:
        nf = fmt          # from the measure spec in load_splitout.py
        for rg in regions:
            for pos in range(1, NC + 1):
                mr = mi_row[(key, rg, pos)]
                ws.cell(row=r, column=1, value=key)
                ws.cell(row=r, column=2, value=rg)
                ws.cell(row=r, column=3, value=pos)
                c = ws.cell(row=r, column=4, value=f'=MAP_SplitIndex!$D{mr}')
                c.number_format = '@'
                c.font = GREEN
                ws.cell(row=r, column=5, value=f'=MAP_SplitIndex!$E{mr}').font = GREEN
                ws.cell(row=r, column=6, value=f'=MAP_SplitIndex!$F{mr}').font = GREEN
                ws.cell(row=r, column=7, value=f'=MAP_SplitIndex!$H{mr}').font = GREEN
                ws.cell(row=r, column=8, value=f'=MAP_SplitIndex!$I{mr}').font = GREEN
                for y in range(nyr):
                    dcol, icol, tcol = CIY + y * 3, CIY + y * 3 + 1, CIY + y * 3 + 2
                    wrow = f'MATCH($B{r}&"|"&{y + 1},{WIDEKEY},0)'
                    ws.cell(row=r, column=tcol, value=(
                        f'=IF($G{r}=0,0,SUMPRODUCT('
                        f'{whole_row(MAT[key], f"$G{r}", NC)},'
                        f'{whole_row(WIDE, wrow, NC)}))')).number_format = nf
                    ws.cell(row=r, column=dcol, value=(
                        f'=IF(OR($G{r}=0,$H{r}=0),0,INDEX({WIDE},{wrow},$C{r})*'
                        f'INDEX({MAT[key]},$H{r},$C{r}))')).number_format = nf
                    ws.cell(row=r, column=icol, value=(
                        f'=${CL(tcol)}{r}-${CL(dcol)}{r}')).number_format = nf
                for off, base in enumerate([0, 1, 2]):
                    terms = '+'.join(f'${CL(CIY + y * 3 + base)}{r}' for y in range(nyr))
                    ws.cell(row=r, column=CIDT + off,
                            value=f'={terms}').number_format = nf
                ci_row[(key, rg, pos)] = r
                r += 1
    CI1 = r - 1
    ws.freeze_panes = 'I6'
    # No autofilter here either: OUT_ByIndustry links to these rows by position.
    ws.cell(row=4, column=1, value='DO NOT SORT OR INSERT ROWS - OUT_ByIndustry and '
                                   'OUT_ByDivision read these rows. Filter on '
                                   'OUT_ByIndustry instead.').font = WARN
    for col, w in zip(['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H'],
                      [11, 8, 9, 10, 50, 9, 9, 9]):
        ws.column_dimensions[col].width = w
    CIM = f'CALC_IndustryImpact!$A${CI0}:$A${CI1}'
    CIR = f'CALC_IndustryImpact!$B${CI0}:$B${CI1}'
    CIDIV = f'CALC_IndustryImpact!$F${CI0}:$F${CI1}'

    # ============================================================ OUT_ByIndustry
    ws = wb.create_sheet('OUT_ByIndustry')
    band(ws, 1, 'OUT_ByIndustry - direct and indirect impact by receiving industry, '
                'all years', 9)
    ws.cell(row=2, column=1, value='=IF(QA_Checks!$C$4="OK","",'
                                   '"CHECK QA_Checks BEFORE USING THESE NUMBERS")'
            ).font = WARN
    note(ws, 3, 'Open effects only - initial plus production-induced. Value added is at '
                'MARKET prices; the master model\'s headline is basic prices. Year detail '
                'is on CALC_IndustryImpact.')
    hdr(ws, 5, ['Measure', 'Region', 'IOIG code', 'Industry', 'Division', 'Division name',
                'Direct', 'Indirect', 'TOTAL'])
    OI0 = 6
    r = OI0
    for key, suffix, label, unit, fmt, mblock in measures:
        nf = fmt          # from the measure spec in load_splitout.py
        for rg in regions:
            for pos in range(1, NC + 1):
                cr = ci_row[(key, rg, pos)]
                ws.cell(row=r, column=1, value=label)
                ws.cell(row=r, column=2, value=rg)
                c = ws.cell(row=r, column=3, value=f'=CALC_IndustryImpact!$D{cr}')
                c.number_format = '@'
                c.font = GREEN
                ws.cell(row=r, column=4,
                        value=f'=CALC_IndustryImpact!$E{cr}').font = GREEN
                ws.cell(row=r, column=5,
                        value=f'=CALC_IndustryImpact!$F{cr}').font = GREEN
                ws.cell(row=r, column=6, value=(
                    f'=IFERROR(INDEX(Lists!$C${DV0}:$C${DV1},'
                    f'MATCH($E{r},Lists!$B${DV0}:$B${DV1},0)),"")')).font = GREEN
                for off in range(3):
                    c = ws.cell(row=r, column=7 + off,
                                value=f'=CALC_IndustryImpact!{CL(CIDT + off)}{cr}')
                    c.number_format = nf
                    c.font = Font(bold=True, color='006100') if off == 2 else GREEN
                r += 1
    OI1 = r - 1
    ws.freeze_panes = 'G6'
    ws.auto_filter.ref = f'A5:F{OI1}'
    for col, w in zip(['A', 'B', 'C', 'D', 'E', 'F'], [30, 8, 10, 50, 9, 40]):
        ws.column_dimensions[col].width = w

    # ============================================================ OUT_ByDivision
    ws = wb.create_sheet('OUT_ByDivision')
    band(ws, 1, 'OUT_ByDivision - impact by summary industry (ANZSIC division)', 12)
    ws.cell(row=2, column=1, value='=IF(QA_Checks!$C$4="OK","",'
                                   '"CHECK QA_Checks BEFORE USING THESE NUMBERS")'
            ).font = WARN
    note(ws, 3, 'ALL REGIONS is the sum of the regional results, each computed with its own '
                'region\'s split-out matrix. It is NOT the national impact of the combined '
                'shock - a state leaks to the rest of Australia. For the national number, '
                'put the lines against Aus. Never add the two together.')
    cols = ['Measure', 'Region', 'Division', 'Division name', 'Key']
    for y in range(nyr):
        cols += [f'Yr{y + 1} Direct', f'Yr{y + 1} Indirect', f'Yr{y + 1} Total']
    cols += ['Direct, all years', 'Indirect, all years', 'TOTAL, all years']
    hdr(ws, 5, cols)
    ODY = 6
    ODT = ODY + nyr * 3
    OD0 = 6
    r = OD0
    div_codes = [dc for dc, dn in div_list]
    for key, suffix, label, unit, fmt, mblock in measures:
        nf = fmt          # from the measure spec in load_splitout.py
        for rg in regions + ['ALL REGIONS']:
            for di, (dc, dn) in enumerate(div_list):
                ws.cell(row=r, column=1, value=label)
                ws.cell(row=r, column=2, value=rg)
                c = ws.cell(row=r, column=3, value=f'=Lists!$B${DV0 + di}')
                c.font = GREEN
                ws.cell(row=r, column=4, value=f'=Lists!$C${DV0 + di}').font = GREEN
                c = ws.cell(row=r, column=5, value=f'="{key}|"&$B{r}&"|"&$C{r}')
                c.number_format = '@'
                for j in range(nyr * 3 + 3):
                    col = ODY + j
                    scol = CIY + j if j < nyr * 3 else CIDT + (j - nyr * 3)
                    crit = (f'{CIM},"{key}",{CIDIV},$C{r}' if rg == 'ALL REGIONS'
                            else f'{CIM},"{key}",{CIR},$B{r},{CIDIV},$C{r}')
                    ws.cell(row=r, column=col, value=(
                        f'=SUMIFS(CALC_IndustryImpact!${CL(scol)}${CI0}:'
                        f'${CL(scol)}${CI1},{crit})')).number_format = nf
                r += 1
    OD1 = r - 1
    ws.freeze_panes = 'F6'
    ws.auto_filter.ref = f'A5:E{OD1}'
    for col, w in zip(['A', 'B', 'C', 'D', 'E'], [30, 13, 9, 42, 22]):
        ws.column_dimensions[col].width = w
    ODKEY = f'OUT_ByDivision!$E${OD0}:$E${OD1}'

    # ============================================================ OUT_Matrix
    ws = wb.create_sheet('OUT_Matrix')
    band(ws, 1, 'OUT_Matrix - the headline: summary industry down, jurisdiction across',
         5 + nreg)
    ws.cell(row=2, column=1, value='=IF(QA_Checks!$C$4="OK","",'
                                   '"CHECK QA_Checks BEFORE USING THESE NUMBERS")'
            ).font = WARN
    note(ws, 3, 'All years summed. Open effects only. ALL REGIONS is the sum of the '
                'regional results and is not the national impact of the same shock.')
    r = 5
    hdrcols = regions + ['ALL REGIONS']
    for key, suffix, label, unit, fmt, mblock in measures:
        nf = fmt          # from the measure spec in load_splitout.py
        for eff, off in [('Direct', 0), ('Indirect', 1), ('TOTAL', 2)]:
            ws.cell(row=r, column=1, value=f'{label}  -  {eff}').font = H2
            hdr(ws, r + 1, ['Division', 'Division name'] + hdrcols)
            for di, (dc, dn) in enumerate(div_list):
                rr = r + 2 + di
                ws.cell(row=rr, column=1, value=f'=Lists!$B${DV0 + di}').font = GREEN
                ws.cell(row=rr, column=2, value=f'=Lists!$C${DV0 + di}').font = GREEN
                for ci2, rg in enumerate(hdrcols):
                    ws.cell(row=rr, column=3 + ci2, value=(
                        f'=IFERROR(INDEX(OUT_ByDivision!${CL(ODT + off)}${OD0}:'
                        f'${CL(ODT + off)}${OD1},MATCH("{key}|{rg}|"&$A{rr},'
                        f'{ODKEY},0)),0)')).number_format = nf
            tr = r + 2 + len(div_list)
            ws.cell(row=tr, column=1, value='Total').font = H3
            for ci2 in range(len(hdrcols)):
                c = ws.cell(row=tr, column=3 + ci2, value=(
                    f'=SUM({CL(3 + ci2)}{r + 2}:{CL(3 + ci2)}{tr - 1})'))
                c.number_format = nf
                c.font = H3
            r = tr + 3
    ws.column_dimensions['A'].width = 9
    ws.column_dimensions['B'].width = 44
    for i in range(len(hdrcols)):
        ws.column_dimensions[CL(3 + i)].width = 14
    ws.freeze_panes = 'C1'

    # ============================================================ QA_Reconcile
    ws = wb.create_sheet('QA_Reconcile')
    band(ws, 1, 'QA_Reconcile - does the industry split sum to the multipliers the main '
                'model uses', 10)
    note(ws, 2, 'The right-hand side is re-derived straight off the Simple multiplier and '
                'Initial effect rows of the SAME RAW block, without touching the '
                'industry matrix. If the split is right the two agree to rounding. This is '
                'the gate for "it should sum and match the same multipliers we use in the '
                'main model".')
    hdr(ws, 5, ['Measure', 'Region', 'Year', 'Split: TOTAL', 'Multiplier: TOTAL',
                'Diff', 'Split: Direct', 'Multiplier: Direct', 'Diff', 'Shock $m'])
    QR0 = 6
    r = QR0
    for key, suffix, label, unit, fmt, mblock in measures:
        nf = fmt          # from the measure spec in load_splitout.py
        for rg in regions:
            for y in range(nyr):
                ws.cell(row=r, column=1, value=key)
                ws.cell(row=r, column=2, value=rg)
                ws.cell(row=r, column=3, value=y + 1)
                wrow = f'MATCH($B{r}&"|"&$C{r},{WIDEKEY},0)'
                tcol, dcol = CIY + y * 3 + 2, CIY + y * 3
                ws.cell(row=r, column=4, value=(
                    f'=SUMIFS(CALC_IndustryImpact!${CL(tcol)}${CI0}:'
                    f'${CL(tcol)}${CI1},{CIM},$A{r},{CIR},$B{r})')).number_format = nf
                simr = f'MATCH($B{r}&"|Simple multiplier",{RK[key]},0)'
                ws.cell(row=r, column=5, value=(
                    f'=SUMPRODUCT({whole_row(MAT[key], simr, NC)},'
                    f'{whole_row(WIDE, wrow, NC)})')).number_format = nf
                ws.cell(row=r, column=6, value=(
                    f'=ROUND($D{r}-$E{r},6)')).number_format = NUM6
                ws.cell(row=r, column=7, value=(
                    f'=SUMIFS(CALC_IndustryImpact!${CL(dcol)}${CI0}:'
                    f'${CL(dcol)}${CI1},{CIM},$A{r},{CIR},$B{r})')).number_format = nf
                inir = f'MATCH($B{r}&"|Initial effect",{RK[key]},0)'
                ws.cell(row=r, column=8, value=(
                    f'=SUMPRODUCT({whole_row(MAT[key], inir, NC)},'
                    f'{whole_row(WIDE, wrow, NC)})')).number_format = nf
                ws.cell(row=r, column=9, value=(
                    f'=ROUND($G{r}-$H{r},6)')).number_format = NUM6
                ws.cell(row=r, column=10, value=(
                    f'=INDEX(CALC_ShockWide!${CL(CWD + NC)}${CW0}:'
                    f'${CL(CWD + NC)}${CW1},{wrow})')).number_format = MONEY
                r += 1
    QR1 = r - 1
    ws.freeze_panes = 'D6'
    for col, w in zip('ABCDEFGHIJ', [11, 8, 6, 16, 18, 12, 16, 18, 12, 12]):
        ws.column_dimensions[col].width = w

    # ============================================================ QA_Checks
    ws = wb.create_sheet('QA_Checks')
    band(ws, 1, 'QA_Checks', 6)
    ws.cell(row=3, column=1, value='Overall').font = H2
    ws.cell(row=4, column=1, value='All gates').font = H3
    QG0 = 8
    gates = [
        ('Both split-out families loaded for every region',
         f'=IF(AND(COUNTIF({raw_tab[measures[0][0]]}!$D$6:$D${raw_last[measures[0][0]]},'
         f'"Spine")={nreg * sum(1 for c, _ in order if not c.startswith("99"))},'
         f'COUNTIF({raw_tab[measures[-1][0]]}!$D$6:$D${raw_last[measures[-1][0]]},'
         f'"Spine")={nreg * sum(1 for c, _ in order if not c.startswith("99"))}),'
         f'"PASS","FAIL")',
         'A region silently missing from the stack, which would quietly zero it'),
        ('Every RAW row MAP landed on carries the code it was looked up by',
         f'=IF(COUNTIF(MAP_SplitIndex!$K${MI0}:$K${MI1},"OK")={MI1 - MI0 + 1},'
         f'"PASS","FAIL")',
         'A misaligned MATCH, which would apply one industry\'s coefficients to another'),
        ('CALC_ShockWide column order matches CALC_ShockByIndustry',
         f'=IF(COUNTIF(CALC_ShockWide!${CL(CWD)}$6:${CL(CWD + NC - 1)}$6,"X")=0,'
         f'"PASS","FAIL")',
         'The layout change silently permuting the spine'),
        ('CALC_ShockWide row totals match the long table',
         f'=IF(SUMPRODUCT(MAX(ABS(CALC_ShockWide!${CL(CWD + NC + 1)}${CW0}:'
         f'${CL(CWD + NC + 1)}${CW1})))<0.000001,"PASS","FAIL")',
         'A row linked to the wrong region or year'),
        ('Every pasted line has a known region',
         f'=IF(COUNTIF(IN_DirectVector!${CL(IVR)}${IV0}:${CL(IVR)}${IV1},'
         f'"UNKNOWN REGION")=0,"PASS","FAIL")',
         'A line that resolves to nothing and is dropped without a trace. Counts '
         'the check column directly, so a bad region shows up even in a year the '
         'line happens to be zero'),
        ('Every pasted code is on the 114-code spine',
         f'=IF(COUNTIF(IN_DirectVector!${CL(IVC)}${IV0}:${CL(IVC)}${IV1},'
         f'"NOT ON THE 114-CODE SPINE")=0,"PASS","FAIL")',
         'A code the split-out tables have no column for'),
        ('Nothing shocked against a dummy row',
         f'=IF(SUMPRODUCT((LEFT(CALC_ShockByIndustry!$B${CS0}:$B${CS1},2)="99")*'
         f'ABS(CALC_ShockByIndustry!${CL(CSTOT)}${CS0}:${CL(CSTOT)}${CS1}))=0,'
         f'"PASS","FAIL")',
         'Spend landing on 9901-9910, which carry no multiplier'),
        ('Shock by industry totals back to the paste',
         f'=IF(ROUND(SUM(CALC_ShockByIndustry!${CL(CSTOT)}${CS0}:${CL(CSTOT)}${CS1})'
         f'-SUM(IN_DirectVector!$D${IV0}:${CL(3 + nyr)}${IV1}),6)=0,"PASS","FAIL")',
         'A line dropped between the paste and the spine'),
        ('Industry split sums to the supplied Simple multiplier',
         f'=IF(SUMPRODUCT(MAX(ABS(QA_Reconcile!$F${QR0}:$F${QR1})))<0.000001,'
         f'"PASS","FAIL")',
         'THE gate: the totals must equal the multipliers the main model uses'),
        ('Direct effect sums to the supplied Initial effect',
         f'=IF(SUMPRODUCT(MAX(ABS(QA_Reconcile!$I${QR0}:$I${QR1})))<0.000001,'
         f'"PASS","FAIL")',
         'The direct/indirect split being drawn in the wrong place'),
        ('No negative indirect effect',
         f'=IF(COUNTIF(CALC_IndustryImpact!${CL(CIDT + 1)}${CI0}:'
         f'${CL(CIDT + 1)}${CI1},"<-0.000001")=0,"PASS","FAIL")',
         'The direct/indirect split drawn in the wrong place. The initial effect '
         'never exceeds the diagonal in any region or measure, so a negative here '
         'means the subtraction has gone wrong'),
        ('Every division on the spine is in the Lists division table',
         f'=IF(COUNTIF(Lists_Spine!$D${SP0}:$D${SP1},"UNMAPPED")=0,"PASS","FAIL")',
         'An IOIG code with no summary industry, which would vanish from OUT_Matrix'),
        ('OUT_ByDivision totals back to OUT_ByIndustry',
         f'=IF(ROUND(SUMIFS(OUT_ByDivision!${CL(ODT + 2)}${OD0}:${CL(ODT + 2)}${OD1},'
         f'OUT_ByDivision!$B${OD0}:$B${OD1},"<>ALL REGIONS")'
         f'-SUM(OUT_ByIndustry!$I${OI0}:$I${OI1}),6)=0,"PASS","FAIL")',
         'A division dropped on the way to the summary'),
    ]
    hdr(ws, 7, ['#', 'Gate', 'Result', 'What it catches'])
    for i, (name, formula, catches) in enumerate(gates):
        r = QG0 + i
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value=name)
        c = ws.cell(row=r, column=3, value=formula)
        c.font = H3
        ws.cell(row=r, column=4, value=catches).font = NOTE
    QG1 = QG0 + len(gates) - 1
    ws.cell(row=4, column=3, value=(
        f'=IF(COUNTIF($C${QG0}:$C${QG1},"PASS")={len(gates)},"OK","NOT OK")')).font = H1
    ws.cell(row=4, column=3).fill = TITLEFILL

    IN0 = QG1 + 3
    ws.cell(row=IN0, column=1, value='For information - not gates').font = H2
    infos = [
        ('Example data still in IN_DirectVector',
         f'=IF(COUNTIF(IN_DirectVector!$A${IV0}:$A${IV1},"line")'
         f'+COUNTIF(IN_DirectVector!$A${IV0}:$A${IV1},"margin")=0,"empty",'
         f'"{len(EXAMPLE_ROWS)} example rows are shipped with this build - replace them")'
         if EXAMPLE else '"built empty"',
         'The shipped example is not a result'),
        ('Value added definition',
         '"MARKET prices (includes taxes on products, P3). The master model headline is '
         'BASIC prices. Different definitions - they will not tie."',
         'The only definition the split-out file supplies'),
        ('Effects included',
         '"Open / Simple multiplier = initial + production-induced. No induced effect."',
         'Closed families are in the source file but not built yet'),
        ('Rows with spend on the paste',
         f'=COUNTIF(IN_DirectVector!${CL(IVT)}${IV0}:${CL(IVT)}${IV1},"<>0")', ''),
        ('Total pasted shock, all years, $m',
         f'=SUM(IN_DirectVector!$D${IV0}:${CL(3 + nyr)}${IV1})', ''),
    ]
    hdr(ws, IN0 + 1, ['#', 'Item', 'Value', 'Note'])
    for i, (name, formula, nt) in enumerate(infos):
        r = IN0 + 2 + i
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=3, value=formula)
        ws.cell(row=r, column=4, value=nt).font = NOTE
    for col, w in zip(['A', 'B', 'C', 'D'], [5, 56, 30, 66]):
        ws.column_dimensions[col].width = w

    # ============================================================ Assumptions
    ws = wb.create_sheet('Assumptions')
    band(ws, 1, 'Assumptions and known limits', 4)
    items = [
        ('Value added is at market prices',
         'The split-out file supplies va@mp only. That includes taxes on products (P3); '
         'the master model\'s headline value added is at basic prices (P1+P2+P4). The two '
         'differ by roughly +0.01 per dollar on average nationally and the sign varies by '
         'industry. Ask the provider for a basic-prices split-out family if the industry '
         'detail has to tie to the headline.'),
        ('Open model only',
         'Simple multiplier - initial plus production-induced. No consumption-induced '
         '(induced) effect, so these numbers are smaller than the master model\'s TOTAL '
         'column and should be compared against Direct + Indirect, not TOTAL.'),
        ('Direct is the initial effect on the shocked industry',
         'The initial effect lands wholly in the industry that received the spend, so '
         'Direct(i) = Shock(i) x InitialEffect(i) and everything else in column i is '
         'indirect. Nothing is assumed about where indirect activity goes - that is what '
         'the split-out matrix supplies.'),
        ('Employment is FTE',
         'Per the state Table 5 column heading in the supplied set. Price year still to '
         'confirm with the provider.'),
        ('Vintage',
         'Split-out tables are FY23 (2022-23), matching the supplied multiplier set. The '
         'ABS margin and tax tables the master uses upstream are 2023-24. Still mixed, '
         'still one year apart, still flagged rather than resolved.'),
        ('6700',
         'No row in the supplied 114-code set. Bridged to 6701 on Lists and again on '
         'IN_DirectVector, so a paste that has not been bridged still resolves.'),
        ('ALL REGIONS',
         'The sum of the regional results, each on its own region\'s matrix. Not the '
         'national impact of the combined shock. Never add it to the Aus row.'),
        ('Summary industries',
         'ANZSIC 2006 divisions A-S, nineteen of them, all represented on the 114-code '
         'spine. Derived from the IOIG subdivision prefix and shown on Lists.'),
    ]
    hdr(ws, 3, ['#', 'Assumption', 'Detail'])
    for i, (a, b) in enumerate(items):
        ws.cell(row=4 + i, column=1, value=i + 1)
        ws.cell(row=4 + i, column=2, value=a).font = H3
        c = ws.cell(row=4 + i, column=3, value=b)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[4 + i].height = 46
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 110

    out = SUBSET_OUT if SUBSET else OUTFILE
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    nf = sum(1 for s in wb.sheetnames for row in wb[s].iter_rows()
             for c in row if isinstance(c.value, str) and c.value.startswith('='))
    print(f'wrote {out.relative_to(ROOT)}')
    print(f'  {len(wb.sheetnames)} tabs, {nf:,} formulas, '
          f'{nreg} regions, {len(measures)} measures, {nyr} years')
    print(f'  IN_DirectVector paste area A{IV0}:{CL(3 + nyr)}{IV1}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
